"""Page-aware text extraction for PDF, TXT, DOCX and PNG/JPG files.

PDF pages with no meaningful selectable text are rendered and read with OCR.
"""

from __future__ import annotations

import io
import logging
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

from PIL import Image

from .config import get_settings
from .ocr import ImageRejected, ocr_image, validate_image

log = logging.getLogger(__name__)

FILE_KINDS = {".pdf": "pdf", ".txt": "txt", ".docx": "docx", ".png": "png", ".jpg": "jpg", ".jpeg": "jpg"}
OCR_RENDER_DPI = 200


class ExtractionError(ValueError):
    """User-facing reason why a file could not be read."""


@dataclass
class PageText:
    page: int
    text: str
    ocr: bool = False


@dataclass
class Extraction:
    pages: list[PageText]
    ocr_required: bool = False
    warnings: list[str] = field(default_factory=list)


def file_kind_of(filename: str) -> str | None:
    return FILE_KINDS.get(Path(filename).suffix.lower())


def validate_upload(filename: str, data: bytes) -> str:
    """Check type, size and basic integrity before anything is stored. Returns the file kind."""
    settings = get_settings()
    kind = file_kind_of(filename)
    if kind is None:
        raise ExtractionError("Unsupported file type. Upload a PDF, DOCX, TXT, PNG or JPG file.")
    if not data:
        raise ExtractionError("This file is empty.")
    if len(data) > settings.max_upload_mb * 1024 * 1024:
        raise ExtractionError(f"This file is larger than {settings.max_upload_mb} MB.")
    if kind == "pdf" and not data.startswith(b"%PDF"):
        raise ExtractionError("This file is not a valid PDF.")
    if kind == "docx" and not data.startswith(b"PK"):
        raise ExtractionError("This file is not a valid DOCX document.")
    if kind == "txt":
        _decode_text(data)
    if kind in ("png", "jpg"):
        try:
            validate_image(data)
        except ImageRejected as exc:
            raise ExtractionError(str(exc)) from None
    return kind


def extract(path: Path, kind: str, on_ocr: Callable[[], None] | None = None) -> Extraction:
    data = path.read_bytes()
    if kind == "pdf":
        result = _extract_pdf(data, on_ocr)
    elif kind == "txt":
        result = Extraction([PageText(i + 1, t) for i, t in enumerate(_decode_text(data).split("\f"))])
    elif kind == "docx":
        result = Extraction([PageText(1, _extract_docx(data))])
    elif kind in ("png", "jpg"):
        if on_ocr:
            on_ocr()
        result = _extract_image(data)
    else:
        raise ExtractionError("Unsupported file type.")

    result.pages = [PageText(p.page, normalize(p.text), p.ocr) for p in result.pages]
    result.pages = [p for p in result.pages if p.text.strip()]
    if not result.pages:
        raise ExtractionError("No readable text was found in this file.")
    return result


def normalize(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)  # also splits ligatures such as "ﬁ"
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return "\n".join(line.strip() for line in text.split("\n")).strip()


def _decode_text(data: bytes) -> str:
    if b"\x00" in data[:4096]:
        raise ExtractionError("This text file appears to be binary, not plain text.")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("This text file could not be decoded. Save it as UTF-8.")


def _extract_pdf(data: bytes, on_ocr: Callable[[], None] | None) -> Extraction:
    import pymupdf

    settings = get_settings()
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception:  # noqa: BLE001
        raise ExtractionError("This PDF is damaged and could not be opened.") from None
    if doc.needs_pass:
        raise ExtractionError("This PDF is password-protected. Upload an unprotected copy.")
    if doc.page_count == 0:
        raise ExtractionError("This PDF has no pages.")

    pages: list[PageText] = []
    scanned: list[int] = []
    for index, page in enumerate(doc):
        text = page.get_text("text", sort=True)
        if sum(ch.isalnum() for ch in text) < settings.scanned_page_min_chars:
            scanned.append(index)
        pages.append(PageText(index + 1, text))

    result = Extraction(pages, ocr_required=bool(scanned))
    if scanned and on_ocr:
        on_ocr()
    for index in scanned:
        pix = doc[index].get_pixmap(dpi=OCR_RENDER_DPI)
        image = Image.open(io.BytesIO(pix.tobytes("png")))
        try:
            ocr = ocr_image(image)
        except Exception as exc:  # noqa: BLE001
            log.warning("OCR failed on page %d: %s", index + 1, exc)
            result.warnings.append(f"Page {index + 1} could not be read with OCR.")
            continue
        if ocr.usable:
            pages[index] = PageText(index + 1, ocr.text, ocr=True)
        elif pages[index].text.strip() or ocr.text.strip():
            pages[index] = PageText(index + 1, ocr.text or pages[index].text, ocr=True)
            result.warnings.append(f"Page {index + 1} has very little readable text.")
    doc.close()
    return result


def _extract_docx(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception:  # noqa: BLE001
        raise ExtractionError("This DOCX file is damaged and could not be opened.") from None
    parts: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        parts.append(f"\n{text}\n" if style.startswith(("heading", "title")) else text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(dict.fromkeys(c for c in cells if c)))
    return "\n".join(parts)


def _extract_image(data: bytes) -> Extraction:
    try:
        image, _ = validate_image(data)
    except ImageRejected as exc:
        raise ExtractionError(str(exc)) from None
    ocr = ocr_image(image)
    if not ocr.usable:
        raise ExtractionError("The image text could not be read. Upload a clearer, higher-resolution image.")
    return Extraction([PageText(1, ocr.text, ocr=True)], ocr_required=True)
