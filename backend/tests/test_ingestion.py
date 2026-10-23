"""Phases 1-4: configuration, ingestion, versions, OCR and restart persistence."""

from __future__ import annotations

import pytest
from conftest import CAMPUS_TEXT, FEES_TEXT, index, make_docx, make_pdf, make_scanned_pdf, text_image

from college_rag import faults
from college_rag.chunking import chunk_pages
from college_rag.config import get_settings
from college_rag.extraction import PageText
from college_rag.ingestion import UploadRejected, reprocess, set_active
from college_rag.retrieval import retrieve
from college_rag.runtime import get_store
from college_rag.store import Store


def test_settings_load_and_storage_created():
    settings = get_settings()
    assert settings.uploads_dir.is_dir() and settings.chroma_dir.is_dir()
    assert settings.top_k > 0 and settings.max_subtasks == 5


def test_pdf_txt_docx_are_indexed_with_page_metadata(college_kb):
    store = get_store()
    for doc_id, version in college_kb.values():
        row = store.get_version(doc_id, version)
        assert row.stage == "indexed" and row.chunk_count >= 1
    assert len(store.searchable_versions()) == 3

    outcome = retrieve("What is the semester 5 tuition fee?")
    assert outcome.method == "vector" and outcome.passages
    top = outcome.passages[0]
    assert top.filename == "fees.pdf" and top.page == 1 and "64,500" in top.text


def test_multi_page_pdf_keeps_page_numbers():
    doc_id, v = index("two.pdf", make_pdf(["Page one is about hostels.", "Page two lists the library hours: 8 am to 9 pm."]),
                      "Two pages")
    outcome = retrieve("library hours")
    assert outcome.passages[0].page == 2


@pytest.mark.parametrize(
    "name,data,message",
    [
        ("notes.exe", b"MZ", "Unsupported file type"),
        ("empty.txt", b"", "empty"),
        ("fake.pdf", b"hello", "not a valid PDF"),
        ("bin.txt", b"\x00\x01\x02", "binary"),
        ("broken.docx", b"not a zip", "not a valid DOCX"),
    ],
)
def test_invalid_uploads_are_rejected(name, data, message):
    with pytest.raises(UploadRejected, match=message):
        index(name, data, "Bad")


def test_oversized_upload_rejected(monkeypatch):
    monkeypatch.setattr(get_settings(), "max_upload_mb", 0)
    with pytest.raises(UploadRejected, match="larger than"):
        index("big.txt", b"x" * 10, "Big")


def test_duplicate_file_is_rejected():
    index("campus.txt", CAMPUS_TEXT.encode(), "Campus")
    with pytest.raises(UploadRejected, match="already uploaded") as exc:
        index("copy.txt", CAMPUS_TEXT.encode(), "Campus copy")
    assert exc.value.status_code == 409


def test_new_version_supersedes_old_one():
    doc_id, _ = index("fees.txt", b"Semester 5 tuition fee is Rs. 50,000.", "Fees", "Fees")
    _, v2 = index("fees-v2.txt", b"Semester 5 tuition fee is Rs. 64,500.", "Fees", "Fees", replaces_id=doc_id)
    assert v2 == 2
    searchable = get_store().searchable_versions()
    assert list(searchable) == [f"{doc_id}:v2"]
    texts = " ".join(p.text for p in retrieve("semester 5 tuition fee").passages)
    assert "64,500" in texts and "50,000" not in texts


def test_failed_new_version_keeps_previous_version_searchable():
    doc_id, _ = index("fees.txt", b"Semester 5 tuition fee is Rs. 64,500.", "Fees", "Fees")
    with faults.injected("embedding"):
        _, v2 = index("fees-v2.txt", b"Semester 5 tuition fee is Rs. 70,000 now.", "Fees", "Fees", replaces_id=doc_id)
    row = get_store().get_version(doc_id, v2)
    assert row.stage == "failed" and row.failed_at == "embedding" and "embedding" in row.error
    assert f"{doc_id}:v1" in get_store().searchable_versions()

    reprocess(doc_id, run_async=False)
    assert get_store().get_version(doc_id, v2).stage == "indexed"
    assert f"{doc_id}:v2" in get_store().searchable_versions()


def test_deactivate_hides_document_and_bumps_generation():
    doc_id, _ = index("campus.txt", CAMPUS_TEXT.encode(), "Campus")
    gen = get_store().kb_generation()
    set_active(doc_id, False)
    assert get_store().kb_generation() == gen + 1
    assert retrieve("where is the campus").passages == []
    set_active(doc_id, True)
    assert retrieve("where is the campus").passages


def test_scanned_pdf_is_read_with_ocr():
    doc_id, v = index("notice-scan.pdf", make_scanned_pdf("Hostel gates close at 9:30 pm every night."), "Hostel notice")
    row = get_store().get_version(doc_id, v)
    assert row.stage == "indexed" and row.ocr_required
    assert "9:30" in retrieve("when do hostel gates close").passages[0].text


def test_png_upload_is_ocr_indexed():
    doc_id, v = index("timetable.png", text_image("Library opens at 8 am on weekdays."), "Library notice")
    assert get_store().get_version(doc_id, v).ocr_required


def test_unreadable_image_upload_fails_safely():
    from PIL import Image
    import io

    buf = io.BytesIO()
    Image.new("RGB", (400, 300), "white").save(buf, "PNG")
    # A valid but blank image passes upload validation; processing must fail cleanly, not crash.
    doc_id, v = index("blank.png", buf.getvalue(), "Blank")
    row = get_store().get_version(doc_id, v)
    assert row.stage == "failed" and row.failed_at == "ocr" and "could not be read" in row.error
    assert get_store().searchable_versions() == {}


def test_restart_persistence(college_kb):
    """A fresh Store/VectorStore on the same storage sees everything, and
    interrupted jobs are marked failed."""
    from college_rag import runtime

    doc_id, v = college_kb["campus"]
    get_store().update_version(doc_id, v, stage="embedding")
    runtime.reset()
    store = Store(get_settings().db_path)
    assert store.mark_interrupted() == 1
    assert store.get_version(doc_id, v).stage == "failed"
    reprocess(doc_id, run_async=False)
    assert retrieve("Where is the campus located?").passages[0].filename == "campus.txt"


def test_chunking_respects_size_overlap_and_pages():
    words = " ".join(f"w{i}" for i in range(1000))
    chunks = chunk_pages([PageText(1, words), PageText(2, "Short second page.")], chunk_tokens=200, overlap_tokens=40)
    page1 = [c for c in chunks if c.page == 1]
    assert len(page1) >= 6
    assert all(len(c.text.split()) <= 150 for c in chunks)
    # Overlap: the next chunk starts with the tail of the previous one.
    assert page1[1].text.split()[0] == page1[0].text.split()[-30]
    assert chunks[-1].page == 2 and chunks[-1].text == "Short second page."


def test_docx_tables_are_extracted():
    import docx, io

    d = docx.Document()
    d.add_paragraph("Office contacts")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Accounts Office", "+91 80 5550 1200"
    buf = io.BytesIO()
    d.save(buf)
    index("contacts.docx", buf.getvalue(), "Contacts")
    assert "5550 1200" in retrieve("accounts office phone").passages[0].text


def test_fees_fixture_text_is_searchable_by_keyword():
    index("fees.pdf", make_pdf([FEES_TEXT]), "Fees", "Fees")
    index("misc.docx", make_docx(["Sports complex opens at 6 am."]), "Sports")
    with faults.injected("chroma"):
        outcome = retrieve("late fee per day")
    assert outcome.method == "bm25" and outcome.passages[0].filename == "fees.pdf"
